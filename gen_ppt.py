"""
生成省赛答辩PPT
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

DARK_GREEN = RGBColor(0x1B, 0x5E, 0x20)
LIGHT_GREEN = RGBColor(0x4C, 0xAF, 0x50)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x21, 0x21, 0x21)
GRAY = RGBColor(0x55, 0x55, 0x55)
LIGHT_GRAY = RGBColor(0xF5, 0xF5, 0xF5)

def add_bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color

def title_slide(title, subtitle='', bg=DARK_GREEN):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, bg)
    tx = slide.shapes.add_textbox(Inches(1), Inches(2.2), Inches(11), Inches(1.5))
    p = tx.text_frame.paragraphs[0]
    p.text = title; p.font.size = Pt(44); p.font.bold = True
    p.font.color.rgb = WHITE; p.alignment = PP_ALIGN.CENTER
    if subtitle:
        tx2 = slide.shapes.add_textbox(Inches(1), Inches(3.8), Inches(11), Inches(1))
        p2 = tx2.text_frame.paragraphs[0]
        p2.text = subtitle; p2.font.size = Pt(20)
        p2.font.color.rgb = RGBColor(0xC8, 0xE6, 0xC9); p2.alignment = PP_ALIGN.CENTER
    return slide

def content_slide(title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, WHITE)
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), prs.slide_width, Inches(1.1))
    bar.fill.solid(); bar.fill.fore_color.rgb = DARK_GREEN
    bar.line.fill.background()
    p = bar.text_frame.paragraphs[0]
    p.text = title; p.font.size = Pt(30); p.font.bold = True
    p.font.color.rgb = WHITE; p.alignment = PP_ALIGN.LEFT
    bar.text_frame.margin_left = Inches(0.8)

    tx = slide.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.5), Inches(5.8))
    tf = tx.text_frame; tf.word_wrap = True
    for i, (txt, lvl) in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = txt
        p.font.size = Pt(20) if lvl == 0 else Pt(16)
        p.font.bold = (lvl == 0)
        p.font.color.rgb = BLACK if lvl == 0 else GRAY
        p.space_after = Pt(12) if lvl == 0 else Pt(3)
    return slide

# ============ 19 SLIDES ============

title_slide('农智云', '基于YOLOv8的农作物病虫害智能诊断系统\n\n深空视界队 · 2026睿抗机器人大赛')

content_slide('目录', [
    ('01  项目背景与痛点分析', 0),
    ('02  系统总体设计', 0),
    ('03  核心技术 — YOLOv8s · 数据增强 · ONNX部署', 0),
    ('04  模型训练与评估报告', 0),
    ('05  系统界面与竞赛验证', 0),
    ('06  大模型智能决策', 0),
    ('07  创新亮点与后续优化', 0),
    ('08  应用前景', 0),
])

content_slide('01 项目背景：农业虫害的严峻挑战', [
    ('全球挑战', 0),
    ('· FAO报告：全球每年因病虫害损失作物产量的10%-28%', 1),
    ('· 粮食安全面临严峻考验，精准防控刻不容缓', 1),
    ('国内现状', 0),
    ('· 年均病虫害发生面积超70亿亩次', 1),
    ('· 传统人工巡查：效率低、覆盖窄、响应滞后，难以应对大面积爆发', 1),
    ('政策机遇', 0),
    ('· 中央一号文件明确促进AI与农业深度融合', 1),
    ('· 智慧农业上升为国家战略，技术赋能正当其时', 1),
])

content_slide('02 痛点分析：传统防治四大瓶颈', [
    ('① 监测效率低 —「人眼巡查」的局限', 0),
    ('   人工巡查覆盖有限，无法大面积高频次实时监测，识别依赖个人经验', 1),
    ('② 诊断能力不足 —「经验判断」的风险', 0),
    ('   病虫害症状相似易混淆，普通农户难以准确判断，错过最佳防治期', 1),
    ('③ 决策支持缺失 —「只看病，不开方」', 0),
    ('   缺乏科学的严重程度量化评估，防治方案为通用模板，无个性化指导', 1),
    ('④ 数据管理空白 —「一次性」管理模式', 0),
    ('   巡查记录难以长期保存追溯，无法支撑趋势预测和主动防控', 1),
])

content_slide('03 系统总体设计：五位一体智能闭环', [
    ('全链路流程', 0),
    ('图像采集 → AI智能检测 → 风险分析 → 大模型防治建议 → 数据归档管理', 1),
    ('模块一：图像采集 — 兼容摄像头、手机、无人机等多终端', 0),
    ('模块二：AI智能检测 — YOLOv8s 精准识别27类农业害虫', 0),
    ('模块三：风险分析 — 数量+置信度+覆盖面积，三维量化风险等级', 0),
    ('模块四：大模型决策 — DeepSeek 生成1500字结构化防治报告', 0),
    ('模块五：数据管理 — 历史检测数据归档、查询、趋势分析', 0),
    ('技术栈', 0),
    ('Python + PyTorch + YOLOv8s + ONNX + PyQt5 + FastAPI + DeepSeek', 1),
])

content_slide('04 核心技术：YOLOv8s 目标检测', [
    ('YOLOv8s — 新一代单阶段目标检测模型（11.2M参数）', 0),
    ('架构特点', 0),
    ('· Anchor-Free：无需预设锚框，简化结构，加速收敛', 1),
    ('· CSPDarknet骨干 + C2f模块：增强多尺度特征提取', 1),
    ('· PAN-FPN颈部：双向多尺度特征融合，兼顾大小目标', 1),
    ('· 解耦检测头：分类与回归独立，精度更高', 1),
    ('关键参数', 0),
    ('· 输入：640×640×3   输出：(1, 31, 8400) — 27类+4坐标', 1),
    ('· 单张推理 4.6ms (RTX3060)  |  i5边缘推理 < 3秒', 1),
    ('· 端到端训练，无缝导出 ONNX，跨平台部署', 1),
    ('为什么选 YOLOv8s？', 0),
    ('· 11.2M参数在精度(76% mAP)与速度(实时)之间达到最佳平衡', 1),
])

content_slide('05 核心技术：摄像头场景数据增强', [
    ('针对比赛平台内置600万工业相机拍摄条件，定制8项增强策略：', 0),
    ('光线与色彩模拟', 0),
    ('· HSV色相 ±3%     →  应对不同色温光源（白炽灯/日光/LED）', 1),
    ('· HSV饱和度±100%  →  应对打印偏色、色彩失真', 1),
    ('· HSV亮度 ±80%    →  应对暗光/强光/背光环境', 1),
    ('拍摄角度模拟', 0),
    ('· 旋转 ±25° + 透视变形  →  纸张倾斜、非正对拍摄', 1),
    ('· 缩放 ±60% + 平移 ±15% →  不同拍摄距离与位置', 1),
    ('鲁棒性增强', 0),
    ('· Mosaic拼图 + MixUp混合 →  复杂背景下的泛化能力', 1),
    ('· 随机遮挡15% →  应对局部反光、对焦不准', 1),
])

content_slide('06 模型优化全路径：四轮迭代', [
    ('从初赛到省赛的持续改进过程：', 0),
    ('', 0),
    ('第一轮：类目精选', 0),
    ('  102类（含仅2张图的弱类）→ 27类精选（每类≥150张训练图）', 1),
    ('  剔除数据不足类别，确保每类都能训出可靠特征', 1),
    ('第二轮：模型升级', 0),
    ('  YOLOv8n (3.2M参数) → YOLOv8s (11.2M参数)', 1),
    ('  480×480 → 640×640，模型容量与输入精度双提升', 1),
    ('第三轮：增强强化', 0),
    ('  基础增强 → 8项摄像头场景专项增强', 1),
    ('第四轮：部署适配', 0),
    ('  PyTorch .pt → ONNX 640，解决比赛平台兼容性问题', 1),
])

content_slide('07 核心技术：ONNX 跨框架部署', [
    ('为什么需要 ONNX？', 0),
    ('· 比赛平台（AI+教学实验平台）运行 ONNX Runtime，不直接支持 .pt', 1),
    ('· ONNX（开放神经网络交换格式）是业界通用推理标准，跨框架跨平台', 1),
    ('部署流程', 0),
    ('· 步骤一：PyTorch + YOLOv8s 训练收敛 → best.pt', 1),
    ('· 步骤二：model.export(format="onnx", imgsz=640, simplify=True)', 1),
    ('· 步骤三：AI+教学平台 → 竞赛评分模块 → 加载ONNX → 实时推理', 1),
    ('关键注意事项', 0),
    ('· 导出尺寸必须为 640×640（480导出的ONNX会导致摄像头模式失败）', 1),
    ('· 本地图片上传正常 ≠ 摄像头模式正常，尺寸不匹配是常见陷阱', 1),
    ('· 其他队YOLOv11 ONNX同方式部署成功，验证了技术路线的可行性', 1),
])

content_slide('08 数据集构建：102类 → 27类精选', [
    ('原始数据：IP102公开数据集，102类农业害虫，16,192张训练图片', 0),
    ('筛选策略', 0),
    ('· 统计每类训练图数量，设定阈值：≥150张 = 保留', 1),
    ('· 结果：筛出27类，覆盖11,876张训练图（原数据的73%）', 1),
    ('最终数据集', 0),
    ('· 训练集：11,876张（平均440张/类，最多2,489张，最少152张）', 1),
    ('· 验证集：2,046张（完全独立隔离，训练过程从未见过）', 1),
    ('· 27类涵盖主要农作物害虫：叶蝉、蚜虫、蝼蛄、金针虫、蛴螬等', 1),
    ('数据集质量保证', 0),
    ('· 训练/验证完全隔离，评估指标真实无偏', 1),
    ('· 标签为YOLO格式（class_id cx cy w h），经Pascal VOC标准验证', 1),
])

content_slide('09 模型评估报告：训练收敛分析', [
    ('训练配置：YOLOv8s | max_epochs=100 | batch=16 | imgsz=640 | patience=20', 0),
    ('实际完成：75/100轮自动收敛终止，无需跑满全部轮次', 0),
    ('收敛证据 → 对应评分标准A档（8-10分）', 0),
    ('· box_loss：1.185 → 1.160（最后10轮仅降0.025，完全水平直线）', 1),
    ('· cls_loss：1.363 → 1.298（缓慢下降后趋于平稳）', 1),
    ('· mAP@0.5：最后10轮波动仅0.006，极度稳定，已达性能瓶颈', 1),
    ('· 结论：模型已充分收敛，挖掘了当前配置下的数据最大潜力 ✓', 1),
    ('核心数值', 0),
    ('· mAP@0.5 = 76.0%  |  mAP@0.5:0.95 = 49.0%', 1),
    ('· Precision = 69.2%  |  Recall = 76.4%', 1),
    ('👉 占位：请插入 runs/train/pest27_final4/results.png', 0),
])

content_slide('10 模型评估报告：各类别性能明细', [
    ('27类细粒度害虫识别，验证集2,046张，各类别表现如下：', 0),
    ('★ mAP > 95%（8类，占30%）', 0),
    ('  mole_cricket 99.5% | Lycorma_delicatula 99.5% | Cicadellidae 98.6%', 1),
    ('  rice_leaf_roller 98.6% | Potosiabre_vitarsis 98.0% | flea_beetle 97.1%', 1),
    ('  grub 96.3% | Ampelophaga 95.1%', 1),
    ('★ mAP 80-95%（6类，占22%）', 0),
    ('  wireworm 87.8% | aphids 87.1% | Cicadella_viridis 86.2%', 1),
    ('  peach_borer 84.7% | Miridae 82.7% | corn_borer 82.3%', 1),
    ('★ mAP 60-80%（6类，占22%）', 0),
    ('  blister_beetle 76.9% | Locustoidea 65.8% | Prodenia_litura 65.4%', 1),
    ('注：27类细粒度识别属复杂视觉任务，76%在该领域属较高水平', 0),
    ('👉 占位：请插入 runs/train/pest27_final4/pest27_final4_val/confusion_matrix_normalized.png', 0),
])

content_slide('11 竞赛验证：20张实测全部成功', [
    ('在AI+教学实验平台使用best.onnx + 内置600万工业相机验证：', 0),
    ('实测结果：20/20 全部检测成功 ✅', 0),
    ('· 覆盖14个不同类别', 1),
    ('· 置信度范围：0.74 - 0.89', 1),
    ('· 推理速度：20张 < 60秒（远低于180秒满分基准）', 1),
    ('· 平台摄像头模式正常调用，无报错无卡顿', 1),
    ('场景多样性', 0),
    ('· 涵盖叶片、茎秆、土壤、果实等多种农业场景', 1),
    ('· 不同角度、光照条件下的稳定检测', 1),
    ('评分预期', 0),
    ('· 实测样本准确度：20/20 × 1.75 = 35分满分', 1),
    ('· 推理时效性：< 180秒 = 15分满分', 1),
    ('👉 占位：请插入 runs/detect/pictures20_result/ 检测效果图', 0),
])

content_slide('12 系统界面：PyQt5 桌面应用', [
    ('完整GUI应用程序（非命令行脚本），满足评分标准A档交互要求：', 0),
    ('左侧控制面板', 0),
    ('· 图片/视频/摄像头三种输入模式自由切换', 1),
    ('· 模型选择 + 置信度阈值 + 设备参数可视化调节', 1),
    ('右侧检测画面', 0),
    ('· 实时显示推理视频流，动态标注害虫位置+类别+置信度', 1),
    ('· 红色检测框 + 白色标签，结果一目了然', 1),
    ('底部结果面板', 0),
    ('· 结构化表格：文件名、类别、置信度、坐标、时间戳', 1),
    ('· 导出格式：TXT / CSV / JSON / Excel 四种', 1),
    ('· 历史回看：缩略图/列表形式直观展示抓拍图片（满足评委复核需求）', 1),
    ('👉 占位：请插入 PyQt5 主界面运行截图', 0),
])

content_slide('13 大模型智能决策：DeepSeek赋能', [
    ('从「发现问题」到「解决问题」的完整闭环', 0),
    ('工作流程', 0),
    ('· 检测结果（虫类+数量+坐标+置信度）→ DeepSeek API → 防治报告', 1),
    ('生成内容（约1500字结构化报告，含以下四大维度）', 0),
    ('· 农业防治：合理轮作、清除杂草、加强水肥管理、增强植株抗性', 1),
    ('· 物理防治：频振式杀虫灯诱杀、防虫网阻隔、人工摘除病株虫卵', 1),
    ('· 生物防治：保护瓢虫/寄生蜂等天敌、科学使用生物农药以虫治虫', 1),
    ('· 化学防治：精准推荐高效低毒低残留环保农药，提供详细施用方案', 1),
    ('技术特性', 0),
    ('· 支持多模型切换：DeepSeek / OpenAI / 智谱GLM / 豆包 / 千帆', 1),
    ('· 本地离线可用（基础知识库），联网时增强（大模型API）', 1),
])

content_slide('14 创新亮点', [
    ('① 数据驱动的类目精选策略', 0),
    ('   基于统计分析筛除数据不足类别，27类精选确保每类训出稳定可靠特征', 1),
    ('② 摄像头场景专项数据增强', 0),
    ('   针对比赛平台工业相机，定制8项增强策略，光线/角度/打印/模糊全覆盖', 1),
    ('③ ONNX跨框架通用部署', 0),
    ('   PyTorch训练 → ONNX导出 → 平台即插即用，突破框架限制', 1),
    ('④ 大模型智能防治决策闭环', 0),
    ('   「识别→分析→建议→存档」全链路闭环，不止于检测', 1),
    ('⑤ 完整桌面级GUI交互系统', 0),
    ('   PyQt5图形界面，三模式输入+四格式导出+历史回看，竞赛交互满分', 1),
    ('⑥ Model Soup多模型融合探索', 0),
    ('   通过权重平均融合不同训练阶段的模型特征，为突破单模型瓶颈提供路径', 1),
])

content_slide('15 后续优化方向', [
    ('精度提升方向', 0),
    ('· Model Soup：多训练轮次权重平均融合（已完成架构验证，待进一步实验）', 1),
    ('· 多模型投票：EfficientNet/ResNet + YOLOv8s 多模型联合推理', 1),
    ('· YOLOv8m：25.9M参数版本，在更强硬件上探索更高精度上限', 1),
    ('推理加速方向', 0),
    ('· TensorRT：ONNX → TensorRT引擎，边缘设备推理速度提升2-3倍', 1),
    ('· OpenVINO：适配Intel平台，i5边缘设备更高效运行', 1),
    ('功能扩展方向', 0),
    ('· 移动端APP（Android/iOS）：田间地头即时拍照识别', 1),
    ('· 无人机巡检集成：大面积农田自动化巡查预警', 1),
    ('· 病害预测模型：基于历史数据训练，从被动诊断升级为主动预测', 1),
])

content_slide('16 应用前景', [
    ('核心落地场景', 0),
    ('· 智慧农场：日常自动化巡检，降低人力成本，提高监测频次与覆盖面积', 1),
    ('· 农技推广站：为基层农技人员提供高效辅助诊断工具，提升服务触达效率', 1),
    ('· 科研教学：为农业院校和科研机构提供实训平台和真实数据支撑', 1),
    ('生态蓝图', 0),
    ('· PaaS开放平台：连接上下游合作伙伴，汇聚农业产业链资源', 1),
    ('· 以数据为核心，以AI为驱动，构建现代农业数字生态系统', 1),
    ('', 0),
    ('愿景：让每一片农田都有 AI 守护', 0),
])

title_slide('感谢聆听', '农智云 — 让每一片农田都有 AI 守护\n\n深空视界队')

# Save
out = 'D:/Dev Projects/Pycharm/yolov8/农智云_省赛PPT.pptx'
prs.save(out)
print(f'✅ PPT已生成: {out}')
print(f'共 {len(prs.slides)} 页')
